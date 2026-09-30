"""Confere a apresentação gerada, sem abrir o PowerPoint.

A apresentação sai de `scripts/gerar-apresentacao.mjs`, e script que posiciona
caixa por coordenada erra calado: o texto que não cabe não estoura no arquivo,
estoura na projeção da sala. Esta conferência é geométrica e feita sobre o
próprio .pptx: forma fora do slide, caixa de texto sobrepondo outra, margem
insuficiente e texto que não cabe na caixa em que foi posto.

    python scripts/validar-apresentacao.py [arquivo.pptx]

A estimativa de texto é aproximada de propósito: ela existe para pegar estouro
grosseiro, e por isso trabalha com folga antes de reprovar.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

PADRAO = Path("PermaneIA-Apresentacao.pptx")
MARGEM_MINIMA = 0.4            # polegadas
SOBREPOSICAO_TOLERADA = 0.02   # polegadas de encosto, que não é colisão

# Largura média de caractere, em fração do corpo da fonte. A Verdana é larga:
# usar o 0,5 de uma Arial aqui deixaria passar estouro de verdade.
LARGURA_POR_FONTE = {"verdana": 0.58, "arial": 0.5, "calibri": 0.48}
LARGURA_PADRAO = 0.55


def pol(valor) -> float:
    return Emu(valor).inches if valor is not None else 0.0


def largura_media(tamanho: float, fonte: str | None) -> float:
    fator = LARGURA_POR_FONTE.get((fonte or "").lower(), LARGURA_PADRAO)
    return tamanho * fator / 72


def cabe(texto: str, larg: float, alt: float, tamanho: float, fonte: str | None) -> bool:
    if not texto.strip():
        return True
    por_linha = max(1, int(larg / largura_media(tamanho, fonte)))
    linhas = 0
    for paragrafo in texto.split("\n"):
        linhas += max(1, -(-len(paragrafo) // por_linha))
    altura_linha = tamanho * 1.35 / 72
    return linhas * altura_linha <= alt + 0.06


def main() -> int:
    arquivo = Path(sys.argv[1]) if len(sys.argv) > 1 else PADRAO
    if not arquivo.exists():
        print(f"REPROVADO: {arquivo} não existe. Rode npm run gerar:apresentacao")
        return 1

    falhas: list[str] = []
    avisos: list[str] = []
    verificacoes = 0

    pres = Presentation(arquivo)
    largura_slide, altura_slide = pol(pres.slide_width), pol(pres.slide_height)
    print(f"  {arquivo.name}: {len(pres.slides)} slides de "
          f"{largura_slide:.2f} x {altura_slide:.2f} polegadas\n")

    for indice, slide in enumerate(pres.slides, 1):
        caixas = []
        mudas: list[tuple[float, float, float, float]] = []
        com_texto = 0

        for forma in slide.shapes:
            x, y = pol(forma.left), pol(forma.top)
            w, h = pol(forma.width), pol(forma.height)

            verificacoes += 1
            if (x < -0.01 or y < -0.01 or x + w > largura_slide + 0.01
                    or y + h > altura_slide + 0.01):
                falhas.append(f"slide {indice}: forma fora do slide em "
                              f"({x:.2f}, {y:.2f}) medindo {w:.2f} x {h:.2f}")

            verificacoes += 1
            if (x < MARGEM_MINIMA - 0.01 or y < 0.2
                    or largura_slide - (x + w) < MARGEM_MINIMA - 0.01):
                avisos.append(f"slide {indice}: forma encostada na borda "
                              f"em ({x:.2f}, {y:.2f})")

            tem_texto = forma.has_text_frame and forma.text_frame.text.strip()
            if not tem_texto:
                if w > 0.05 and h > 0.05:
                    mudas.append((x, y, w, h))
                continue

            com_texto += 1
            texto = forma.text_frame.text
            corpos, fontes = [], []
            for paragrafo in forma.text_frame.paragraphs:
                for run in paragrafo.runs:
                    if run.font.size:
                        corpos.append(run.font.size.pt)
                    if run.font.name:
                        fontes.append(run.font.name)
            tamanho = max(corpos) if corpos else 12
            fonte = fontes[0] if fontes else None
            verificacoes += 1
            if not cabe(texto, w, h, tamanho, fonte):
                falhas.append(f"slide {indice}: texto estoura a caixa de "
                              f"{w:.2f} x {h:.2f} em {tamanho:.1f} pt "
                              f"({len(texto)} caracteres): "
                              f"{texto[:40].replace(chr(10), ' ')!r}")
            caixas.append((x, y, w, h, texto[:28].replace("\n", " ")))

        # Sobreposição entre caixas de texto. Texto dentro de forma sem texto é
        # card, e card é desenho normal, não colisão.
        for i in range(len(caixas)):
            for j in range(i + 1, len(caixas)):
                ax, ay, aw, ah, at = caixas[i]
                bx, by, bw, bh, bt = caixas[j]
                sx = min(ax + aw, bx + bw) - max(ax, bx)
                sy = min(ay + ah, by + bh) - max(ay, by)
                verificacoes += 1
                if sx > SOBREPOSICAO_TOLERADA and sy > SOBREPOSICAO_TOLERADA:
                    falhas.append(f"slide {indice}: {at!r} e {bt!r} se sobrepõem "
                                  f"em {sx:.2f} x {sy:.2f} polegadas")

        notas = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notas = slide.notes_slide.notes_text_frame.text.strip()
        print(f"  slide {indice:2d}: {len(slide.shapes):2d} formas, {com_texto:2d} com "
              f"texto{'  · com nota do apresentador' if notas else ''}")

    print()
    if avisos:
        print(f"AVISOS ({len(avisos)}), que merecem olhada mas não reprovam:")
        for aviso in avisos:
            print(f"  - {aviso}")
        print()
    if falhas:
        print(f"REPROVADO: {len(falhas)} problema(s) em {verificacoes} verificações")
        for falha in falhas:
            print(f"  - {falha}")
        return 1
    print(f"APROVADO: {verificacoes} verificações de geometria, nenhuma falha")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
